import openpyxl
import json
import os
from datetime import datetime
from dotenv import load_dotenv

load_dotenv()

load_dotenv()

from app import app
from utils.db import Session

file_path = r'C:\Users\azlon\Downloads\Report-ISP -LA\SOC September 2026.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

print('[TEST] Testing SOC import with actual file data\n')

templates = {
    'Raw_Threat_HC': 'threat_hc',
    'Raw_URL_Blocked': 'url_blocked',
    'Raw_Traffic_Rule': 'traffic_rule',
    'Raw_Traffic_Daily': 'traffic_daily'
}

test_results = []

for sheet_name, template_type in templates.items():
    ws = wb[sheet_name]
    headers = [str(cell.value) for cell in ws[1] if cell.value]
    
    rows_data = []
    for i in range(2, min(12, ws.max_row + 1)):
        row_dict = {}
        for j, header in enumerate(headers):
            cell = ws.cell(row=i, column=j+1)
            val = cell.value
            if isinstance(val, datetime):
                val = val.strftime('%Y-%m-%d %H:%M:%S')
            row_dict[header] = val
        rows_data.append(row_dict)
    
    print('Testing: ' + template_type)
    print('  Rows to import: ' + str(len(rows_data)))
    
    with app.test_client() as client:
        with client.session_transaction() as sess:
            sess['user_id'] = '8f3d5499-3d49-48c1-b496-299bb6d917d0'
            sess['user_role'] = 'Admin'
        
        detect_res = client.post('/api/import/soc/detect',
            json={'headers': headers},
            content_type='application/json'
        )
        detect_data = json.loads(detect_res.data)
        
        conf = detect_data['data']['confidence']
        print('  Detected: ' + detect_data['data']['template_type'] + ' (confidence: {:.1%})'.format(conf))
        
        exec_res = client.post('/api/import/soc/execute',
            json={
                'data': rows_data,
                'template_type': template_type,
                'replace_existing': False,
                'period_start': '2026-09-01',
                'period_end': '2026-09-30'
            },
            content_type='application/json'
        )
        exec_data = json.loads(exec_res.data)
        
        if exec_data['success']:
            inserted = exec_data['data']['inserted']
            skipped = exec_data['data']['skipped']
            print('  Result: {} inserted, {} skipped'.format(inserted, skipped))
            test_results.append({'template': template_type, 'inserted': inserted, 'skipped': skipped})
        else:
            print('  ERROR: ' + exec_data.get('error', 'Unknown error'))
    
    print()

wb.close()

print('[VERIFY] Checking database records...\n')
session = Session()
try:
    from sqlalchemy import text
    
    result = session.execute(text('SELECT COUNT(*) as cnt FROM soc_events')).fetchone()
    total_events = result[0] if result else 0
    
    result = session.execute(text('SELECT COUNT(*) as cnt FROM soc_upload_batch')).fetchone()
    total_batches = result[0] if result else 0
    
    print('Total SOC events in DB: ' + str(total_events))
    print('Total upload batches: ' + str(total_batches))
    
    result = session.execute(text('SELECT event_type, COUNT(*) as cnt FROM soc_events GROUP BY event_type')).fetchall()
    print('\nEvents by type:')
    for row in result:
        print('  {}: {}'.format(row[0], row[1]))
    
finally:
    session.close()

print('\n[OK] Import test completed')
