import os
from dotenv import load_dotenv
load_dotenv()
import openpyxl
from routes.api_import import map_row_to_event, SOC_TEMPLATES

file_path = r'C:\Users\azlon\Downloads\Report-ISP -LA\SOC September 2026.xlsx'
wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
ws = wb['Raw_Traffic_Daily']

headers = [str(cell.value) for cell in ws[1] if cell.value]

print('Checking traffic_daily rows from actual file:\n')

for i in range(2, min(12, ws.max_row + 1)):
    row_dict = {}
    for j, header in enumerate(headers):
        cell = ws.cell(row=i, column=j+1)
        row_dict[header] = cell.value
    
    mapped = map_row_to_event(row_dict, 'traffic_daily')
    
    print('Row ' + str(i) + ':')
    if mapped is None:
        print('  ERROR: map_row_to_event returned None')
    else:
        config = SOC_TEMPLATES['traffic_daily']
        is_valid = True
        for f in config['required_fields']:
            val = mapped.get(f)
            is_valid = is_valid and (val is not None)
            print('  ' + f + ': ' + str(val))
        print('  VALID: ' + str(is_valid))
    print()

wb.close()
